"""
Resume Parser Service
Extracts text and skills from PDF and DOCX files.
"""
import re
import json
from pathlib import Path

# Skills dictionary for extraction
KNOWN_SKILLS = {
    # Programming Languages
    'Programming': [
        'Python', 'Java', 'JavaScript', 'TypeScript', 'C', 'C++', 'C#', 'Ruby', 'Go', 'Golang',
        'Rust', 'Swift', 'Kotlin', 'PHP', 'R', 'MATLAB', 'Scala', 'Perl', 'Shell', 'Bash',
        'PowerShell', 'Assembly', 'COBOL', 'Fortran', 'Haskell', 'Erlang', 'Elixir',
        'Dart', 'Groovy', 'Lua', 'Julia'
    ],
    # Web Frameworks
    'Framework': [
        'Flask', 'Django', 'FastAPI', 'React', 'Angular', 'Vue', 'Vue.js', 'Next.js', 'Nuxt.js',
        'Express', 'Node.js', 'Spring Boot', 'Spring', 'Laravel', 'Ruby on Rails', 'Rails',
        'ASP.NET', '.NET', 'Hibernate', 'Struts', 'Play', 'Gin', 'Echo', 'Fiber',
        'Svelte', 'Remix', 'Gatsby', 'Bootstrap', 'Tailwind', 'jQuery', 'Redux',
        'GraphQL', 'REST API', 'gRPC', 'Microservices'
    ],
    # Databases
    'Database': [
        'MySQL', 'PostgreSQL', 'MongoDB', 'SQLite', 'Oracle', 'SQL Server', 'MSSQL',
        'Redis', 'Cassandra', 'Elasticsearch', 'DynamoDB', 'Firebase', 'Firestore',
        'CouchDB', 'Neo4j', 'InfluxDB', 'MariaDB', 'Supabase', 'SQL', 'NoSQL',
        'HBase', 'Memcached'
    ],
    # Cloud
    'Cloud': [
        'AWS', 'Azure', 'GCP', 'Google Cloud', 'Heroku', 'DigitalOcean', 'Linode',
        'CloudFlare', 'Vercel', 'Netlify', 'IBM Cloud', 'Oracle Cloud',
        'EC2', 'S3', 'Lambda', 'CloudFront', 'RDS', 'EKS', 'ECS', 'Fargate'
    ],
    # DevOps & Tools
    'DevOps': [
        'Docker', 'Kubernetes', 'Jenkins', 'GitHub Actions', 'GitLab CI', 'CircleCI',
        'Terraform', 'Ansible', 'Chef', 'Puppet', 'Helm', 'Prometheus', 'Grafana',
        'Nginx', 'Apache', 'Linux', 'Ubuntu', 'CentOS', 'RHEL', 'Vagrant',
        'Travis CI', 'ArgoCD', 'Istio', 'Kafka', 'RabbitMQ', 'Airflow'
    ],
    # Tools
    'Tool': [
        'Git', 'GitHub', 'GitLab', 'Bitbucket', 'Jira', 'Confluence', 'Slack',
        'VS Code', 'PyCharm', 'IntelliJ', 'Eclipse', 'Postman', 'Swagger',
        'Jupyter', 'Anaconda', 'npm', 'pip', 'Maven', 'Gradle', 'Webpack',
        'Figma', 'Adobe XD', 'Photoshop', 'Linux', 'Vim', 'Bash', 'PowerShell'
    ],
    # AI/ML
    'Other': [
        'Machine Learning', 'Deep Learning', 'TensorFlow', 'PyTorch', 'Keras', 'Scikit-learn',
        'NLP', 'Computer Vision', 'Data Science', 'Pandas', 'NumPy', 'Matplotlib',
        'Seaborn', 'OpenCV', 'NLTK', 'SpaCy', 'Transformers', 'HuggingFace',
        'Data Structures', 'Algorithms', 'System Design', 'OOP', 'Design Patterns',
        'Agile', 'Scrum', 'Kanban', 'SDLC', 'TDD', 'BDD', 'CI/CD',
        'HTML', 'CSS', 'SASS', 'LESS', 'WebSockets', 'OAuth', 'JWT'
    ]
}


def extract_text_from_pdf(file_path):
    """Extract text from a PDF file. Tries pypdf first, falls back to PyPDF2."""
    # Try pypdf (modern, preferred)
    try:
        from pypdf import PdfReader
        text = []
        with open(file_path, 'rb') as f:
            reader = PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
        return '\n'.join(text)
    except ImportError:
        pass
    except Exception as e:
        return f"[Error extracting PDF text: {str(e)}]"
    # Fallback: PyPDF2 (legacy)
    try:
        import PyPDF2
        text = []
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
        return '\n'.join(text)
    except Exception as e:
        return f"[Error extracting PDF text: {str(e)}]"


def extract_text_from_docx(file_path):
    """Extract text from a DOCX file."""
    try:
        from docx import Document
        doc = Document(file_path)
        text = []
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text)
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)
        return '\n'.join(text)
    except Exception as e:
        return f"[Error extracting DOCX text: {str(e)}]"


def extract_text(file_path, file_type):
    """Extract text from resume file based on type."""
    file_type = file_type.lower().strip('.')
    if file_type == 'pdf':
        return extract_text_from_pdf(file_path)
    elif file_type in ('docx', 'doc'):
        return extract_text_from_docx(file_path)
    return ""


def extract_skills_from_text(text):
    """Extract skills from resume text using pattern matching."""
    if not text:
        return []
    
    found_skills = []
    text_lower = text.lower()
    
    for category, skills in KNOWN_SKILLS.items():
        for skill in skills:
            # Use word boundary matching
            pattern = r'\b' + re.escape(skill.lower()) + r'\b'
            if re.search(pattern, text_lower):
                found_skills.append({
                    'name': skill,
                    'category': category
                })
    
    # Deduplicate
    seen = set()
    unique_skills = []
    for s in found_skills:
        if s['name'].lower() not in seen:
            seen.add(s['name'].lower())
            unique_skills.append(s)
    
    return unique_skills


def extract_email(text):
    """Extract email from text."""
    pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    matches = re.findall(pattern, text)
    return matches[0] if matches else None


def extract_phone(text):
    """Extract phone number from text."""
    pattern = r'[\+]?[(]?[0-9]{3}[)]?[-\s\.]?[0-9]{3}[-\s\.]?[0-9]{4,6}'
    matches = re.findall(pattern, text)
    return matches[0] if matches else None


def parse_resume(file_path, file_type):
    """
    Main entry point for resume parsing.
    Returns dict with extracted_text and skills.
    """
    text = extract_text(file_path, file_type)
    skills = extract_skills_from_text(text)
    skill_names = [s['name'] for s in skills]
    
    return {
        'text': text,
        'skills': skills,
        'skill_names': skill_names,
        'email': extract_email(text),
        'phone': extract_phone(text),
        'word_count': len(text.split()) if text else 0
    }
